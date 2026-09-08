"""Extração de texto e divisão em trechos (SPEC §9.2).

PDF: PyMuPDF (fallback pypdf) · DOCX: python-docx · TXT/MD: leitura direta.
Trechos de ~500 tokens (aprox. 4 caracteres/token) com sobreposição de 80 tokens e página quando houver.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

FORMATOS = {".pdf", ".docx", ".txt", ".md"}
TOKENS_TRECHO = 500
TOKENS_SOBREPOSICAO = 80
CHARS_POR_TOKEN = 4


class ArquivoInvalido(Exception):
    pass


class SemTexto(Exception):
    pass


@dataclass
class Pagina:
    numero: Optional[int]
    texto: str


@dataclass
class Extracao:
    paginas: List[Pagina]
    metadados: Dict[str, Any] = field(default_factory=dict)
    n_paginas: int = 0

    @property
    def texto(self) -> str:
        return "\n\n".join(p.texto for p in self.paginas)


def formato_valido(nome: str) -> bool:
    return Path(nome).suffix.lower() in FORMATOS


def extrair(caminho: Path) -> Extracao:
    ext = Path(caminho).suffix.lower()
    if ext == ".pdf":
        return _extrair_pdf(caminho)
    if ext == ".docx":
        return _extrair_docx(caminho)
    if ext in (".txt", ".md"):
        return _extrair_texto(caminho)
    raise ArquivoInvalido(f"Formato não suportado: {ext}")


def _limpar(t: str) -> str:
    t = t.replace("\x00", "")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _extrair_pdf(caminho: Path) -> Extracao:
    paginas: List[Pagina] = []
    meta: Dict[str, Any] = {}
    n = 0
    try:
        import pymupdf  # type: ignore

        with pymupdf.open(str(caminho)) as doc:
            n = doc.page_count
            md = doc.metadata or {}
            meta = {"titulo": (md.get("title") or "").strip(), "autor": (md.get("author") or "").strip(), "ano": "", "ano_arquivo": _ano(md.get("creationDate") or "")}
            for i, pag in enumerate(doc):
                paginas.append(Pagina(numero=i + 1, texto=_limpar(pag.get_text("text") or "")))
    except Exception as e_fitz:  # noqa: BLE001
        try:
            from pypdf import PdfReader

            with open(caminho, "rb") as fh:  # fecha o handle mesmo em falha (Windows)
                reader = PdfReader(fh)
                n = len(reader.pages)
                md = reader.metadata or {}
                meta = {"titulo": (getattr(md, "title", "") or "").strip(), "autor": (getattr(md, "author", "") or "").strip(), "ano": "", "ano_arquivo": ""}
                for i, pag in enumerate(reader.pages):
                    paginas.append(Pagina(numero=i + 1, texto=_limpar(pag.extract_text() or "")))
        except Exception as e_pypdf:  # noqa: BLE001
            raise ArquivoInvalido(f"PDF ilegível ({e_fitz.__class__.__name__}/{e_pypdf.__class__.__name__})") from e_pypdf
    if n == 0:
        raise ArquivoInvalido("PDF sem páginas")
    total = sum(len(p.texto) for p in paginas)
    if total < 20:
        raise SemTexto("PDF sem texto extraível")
    return Extracao(paginas=paginas, metadados=meta, n_paginas=n)


def _extrair_docx(caminho: Path) -> Extracao:
    try:
        import docx  # type: ignore

        d = docx.Document(str(caminho))
    except Exception as e:  # noqa: BLE001
        raise ArquivoInvalido(f"DOCX ilegível ({e.__class__.__name__})") from e
    partes = [p.text for p in d.paragraphs if p.text.strip()]
    for tabela in d.tables:
        for linha in tabela.rows:
            partes.append(" | ".join(c.text.strip() for c in linha.cells if c.text.strip()))
    texto = _limpar("\n".join(partes))
    if len(texto) < 20:
        raise SemTexto("DOCX sem texto")
    cp = d.core_properties
    meta = {"titulo": (cp.title or "").strip(), "autor": (cp.author or "").strip(), "ano": str(cp.created.year) if cp.created else ""}
    return Extracao(paginas=[Pagina(numero=None, texto=texto)], metadados=meta, n_paginas=0)


def _extrair_texto(caminho: Path) -> Extracao:
    try:
        raw = Path(caminho).read_bytes()
    except OSError as e:
        raise ArquivoInvalido(str(e)) from e
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            texto = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ArquivoInvalido("codificação desconhecida")
    texto = _limpar(texto)
    if len(texto) < 20:
        raise SemTexto("arquivo de texto vazio")
    titulo = ""
    m = re.search(r"^#\s+(.+)$", texto, flags=re.M)
    if m:
        titulo = m.group(1).strip()
    return Extracao(paginas=[Pagina(numero=None, texto=texto)], metadados={"titulo": titulo, "autor": "", "ano": ""}, n_paginas=0)


def _ano(data_pdf: str) -> str:
    m = re.search(r"(\d{4})", data_pdf or "")
    return m.group(1) if m else ""


def dividir_em_trechos(extracao: Extracao, tokens: int = TOKENS_TRECHO, sobreposicao: int = TOKENS_SOBREPOSICAO) -> List[Dict[str, Any]]:
    """[{'id': 'sha-000', 'pagina': 3, 'texto': '...'}] — o id é preenchido pelo ingest."""
    alvo = tokens * CHARS_POR_TOKEN
    passo = max(alvo - sobreposicao * CHARS_POR_TOKEN, alvo // 2)
    trechos: List[Dict[str, Any]] = []
    for pag in extracao.paginas:
        texto = pag.texto
        if not texto:
            continue
        inicio = 0
        while inicio < len(texto):
            fim = min(len(texto), inicio + alvo)
            # tenta cortar em fim de frase/parágrafo
            if fim < len(texto):
                corte = max(texto.rfind("\n\n", inicio, fim), texto.rfind(". ", inicio, fim))
                if corte > inicio + alvo // 2:
                    fim = corte + 1
            pedaco = texto[inicio:fim].strip()
            if pedaco:
                trechos.append({"pagina": pag.numero, "texto": pedaco})
            if fim >= len(texto):
                break
            inicio = max(fim - sobreposicao * CHARS_POR_TOKEN, inicio + passo)
    return trechos


def detectar_metadados(extracao: Extracao, nome_original: str) -> Dict[str, Any]:
    """Metadados do arquivo + primeira página (SPEC §9.3). O professor corrige depois."""
    meta = dict(extracao.metadados)
    primeira = extracao.paginas[0].texto if extracao.paginas else ""
    if not meta.get("titulo"):
        linhas = [l.strip() for l in primeira.split("\n") if l.strip()]
        cand = next((l for l in linhas[:8] if 8 <= len(l) <= 120 and not l.isupper() or (l.isupper() and 8 <= len(l) <= 80)), "")
        meta["titulo"] = cand or Path(nome_original).stem.replace("_", " ").replace("-", " ")
    if not meta.get("ano"):
        # ano no texto da primeira página vale mais que a data de criação do arquivo (muitas vezes é a data do download)
        m = re.search(r"\b(19[5-9]\d|20[0-4]\d)\b", primeira[:3000])
        meta["ano"] = m.group(1) if m else (meta.get("ano_arquivo") or "")
    meta.pop("ano_arquivo", None)
    meta["categoria"] = _categoria(primeira + " " + nome_original)
    meta["instrumento"] = _instrumento(primeira + " " + nome_original)
    return meta


_CATEGORIAS = [
    ("pedagogia musical", ["pedagogia musical", "educação musical", "music education", "ensino de música"]),
    ("teoria", ["teoria musical", "music theory", "escala", "intervalo"]),
    ("percepção", ["percepção", "ear training", "solfejo", "ditado"]),
    ("harmonia", ["harmonia", "harmony", "acorde", "encadeamento"]),
    ("história", ["história da música", "music history", "período", "barroco", "romantismo"]),
    ("técnica instrumental", ["técnica", "technique", "dedilhado", "postura", "embocadura"]),
    ("repertório", ["repertório", "repertoire", "partitura", "peças"]),
    ("educação infantil", ["educação infantil", "musicalização", "crianças pequenas", "early childhood"]),
    ("avaliação", ["avaliação", "assessment", "rubrica"]),
    ("currículo", ["currículo", "bncc", "habilidade ef", "curricular"]),
    ("inclusão", ["inclusão", "inclusiva", "deficiência", "tea", "autismo"]),
    ("metodologia", ["orff", "kodály", "kodaly", "dalcroze", "gordon", "swanwick", "suzuki"]),
]


def _categoria(texto: str) -> str:
    t = texto.lower()
    melhor, pontos = "outros", 0
    for cat, chaves in _CATEGORIAS:
        p = sum(t.count(k) for k in chaves)
        if p > pontos:
            melhor, pontos = cat, p
    return melhor


def _instrumento(texto: str) -> str:
    t = texto.lower()
    for nome, ident in [("piano", "piano"), ("teclado", "teclado"), ("violão", "violao"), ("guitarra", "guitarra"), ("flauta doce", "flauta_doce"), ("flauta transversal", "flauta_transversal"), ("clarinete", "clarinete"), ("saxofone", "saxofone"), ("trompete", "trompete"), ("trombone", "trombone"), ("violino", "violino"), ("violoncelo", "violoncelo"), ("viola", "viola"), ("canto", "canto"), ("coral", "coral"), ("bateria", "bateria"), ("percussão", "percussao"), ("percepção", "percepcao_musical"), ("musicalização", "musicalizacao")]:
        if nome in t:
            return ident
    return ""

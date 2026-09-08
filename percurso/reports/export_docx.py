"""Exportação DOCX com python-docx (SPEC §12.6). Gráficos PNG embutidos."""
from __future__ import annotations

from pathlib import Path

from .model import Relatorio


def exportar(rel: Relatorio, destino: Path) -> Path:
    import docx
    from docx.shared import Inches, Pt

    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    estilo = d.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    d.core_properties.title = rel.titulo
    d.core_properties.author = "Percurso"
    d.add_heading(rel.titulo, 0)
    sub = " · ".join(x for x in [rel.subtitulo, rel.identificacao, f"código {rel.codigo}" if rel.codigo else "", rel.periodo, f"gerado em {rel.gerado_em}" if rel.gerado_em else ""] if x)
    if sub:
        d.add_paragraph(sub).italic = True
    if rel.professor or rel.instituicao:
        d.add_paragraph(" · ".join(x for x in [rel.professor, rel.instituicao] if x))
    if rel.rotulo_ia:
        p = d.add_paragraph(rel.rotulo_ia)
        p.runs[0].italic = True
    if rel.resumo_numeros:
        t = d.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        for k, v in rel.resumo_numeros.items():
            c = t.add_row().cells
            c[0].text, c[1].text = k, v
        d.add_paragraph("")
    for s in rel.secoes:
        d.add_heading(s.titulo, level=1)
        for p in s.paragrafos:
            d.add_paragraph(p)
        for i in s.itens:
            d.add_paragraph(i, style="List Bullet")
        for tb in s.tabelas:
            d.add_paragraph(tb.titulo).runs[0].bold = True
            t = d.add_table(rows=1, cols=len(tb.cabecalho))
            t.style = "Light Grid Accent 1"
            for j, c in enumerate(tb.cabecalho):
                t.rows[0].cells[j].text = c
            for l in tb.linhas:
                cells = t.add_row().cells
                for j, c in enumerate(l[: len(tb.cabecalho)]):
                    cells[j].text = str(c)
            d.add_paragraph("")
        for g in s.graficos:
            try:
                d.add_picture(g.caminho_png, width=Inches(6.0))
                d.add_paragraph(g.titulo).runs[0].italic = True
            except Exception:  # noqa: BLE001
                d.add_paragraph(f"[gráfico indisponível: {g.titulo}]")
    d.add_paragraph("Percurso — Plataforma de inteligência pedagógica").runs[0].italic = True
    d.save(str(destino))
    return destino

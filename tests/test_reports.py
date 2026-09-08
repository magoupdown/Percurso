"""Relatórios e exportação (casos 12–14): PDF e DOCX abrem e contêm os dados; gráficos; materiais; MIDI/MusicXML."""
from __future__ import annotations

from pathlib import Path

from percurso.core import planner
from percurso.reports import builder, materials
from percurso.ui import texts as T

from mocks.gemini import ClienteFalso


def _texto_pdf(p: Path) -> str:
    import pymupdf

    with pymupdf.open(str(p)) as d:
        return "\n".join(pg.get_text() for pg in d)


def _texto_docx(p: Path) -> str:
    import docx

    d = docx.Document(str(p))
    partes = [x.text for x in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            partes.extend(c.text for c in r.cells)
    return "\n".join(partes)


def test_todos_os_tipos_e_exportacoes(sessao, registro_com_14_aulas, tmp_path):
    reg = registro_com_14_aulas
    sessao.carregar(reg.codigo)
    for tipo in builder.TIPOS:
        rel = builder.montar(sessao, tipo, com_ia=False)
        assert rel.titulo and rel.secoes and rel.codigo == reg.codigo
        saida = builder.exportar(rel, tmp_path, f"r_{tipo}", ["pdf", "docx", "md", "csv"])
        pdf, docx_ = Path(saida["pdf"][0]), Path(saida["docx"][0])
        assert pdf.exists() and pdf.stat().st_size > 1000 and docx_.exists()
        tp, td = _texto_pdf(pdf), _texto_docx(docx_)
        assert reg.codigo in tp and reg.codigo in td
        assert Path(saida["md"][0]).read_text(encoding="utf-8").startswith("# ")
    rel = builder.montar(sessao, "frequencia", com_ia=False)
    assert rel.resumo_numeros["Aulas registradas"] == "14" and rel.resumo_numeros["Canceladas (instituição)"] == "1"
    assert rel.resumo_numeros["Aulas previstas"] == "13" and rel.resumo_numeros["Faltas"] == "1"
    assert rel.resumo_numeros["Frequência"] == f"{round(100 * 12 / 13, 1):.1f} %"
    assert rel.secoes[0].graficos and all(Path(g.caminho_png).exists() for g in rel.secoes[0].graficos)
    tp = " ".join(_texto_pdf(Path(builder.exportar(rel, tmp_path, "freq", ["pdf"])["pdf"][0])).split())
    assert "Frequência" in tp and "Cancelada (instituição)" in tp
    csvs = builder.exportar(rel, tmp_path, "freq", ["csv"])["csv"]
    assert csvs and "Situação" in Path(csvs[0]).read_text(encoding="utf-8-sig")


def test_relatorio_periodo_e_rendimento(sessao, registro_com_14_aulas):
    reg = registro_com_14_aulas
    sessao.carregar(reg.codigo)
    rel = builder.montar(sessao, "rendimento", inicio="2026-06-01", fim="2026-06-30", com_ia=False)
    assert "junho" not in rel.periodo and "01/06/2026" in rel.periodo
    assert rel.resumo_numeros and rel.secoes[0].tabelas[0].cabecalho[0] == "Data"
    assert len(rel.secoes[0].tabelas[0].linhas) == 4  # 5 aulas em junho, 1 delas é falta sem rendimento
    assert rel.secoes[0].graficos
    vazio = builder.montar(sessao, "rendimento", inicio="2030-01-01", com_ia=False)
    assert "Nenhum rendimento" in vazio.secoes[0].paragrafos[0]


def test_pedagogico_narrativo_com_gemini_mock(sessao, registro_com_14_aulas):
    reg = registro_com_14_aulas
    sessao.carregar(reg.codigo)
    sessao.cliente_gemini = ClienteFalso([f"O aluno {reg.codigo} avançou na pulsação e ainda acelera no final das frases."])
    sessao.modo_ia, sessao.gemini_pronto = "gemini", True
    rel = builder.montar(sessao, "pedagogico", com_ia=True)
    assert rel.rotulo_ia and rel.secoes[0].titulo == "Síntese narrativa"
    assert rel.secoes[0].paragrafos[0].startswith("João avançou")
    assert "joão" not in sessao.cliente_gemini.prompts[0].lower()
    assert any(s.titulo == "Recomendações" for s in rel.secoes)


def test_materiais_e_exemplos_musicais(sessao, registro_com_14_aulas, tmp_path):
    reg = registro_com_14_aulas
    sessao.carregar(reg.codigo)
    plano = planner.gerar_plano(sessao.entrada_planejamento())
    rubricas = sessao.repo.ler_rubricas()
    rels = [
        materials.folha_professor(reg, plano, None, rubricas),
        materials.folha_aluno(reg, plano, None),
        materials.exercicios(reg, plano, None),
        materials.rubrica(reg, plano.criterios, None, rubricas),
        materials.ficha(reg, sessao.contexto.estado, None, "Piano"),
        materials.lista_repertorio(reg, sessao.contexto.repertorio, None),
    ]
    for i, r in enumerate(rels):
        saida = materials.salvar(r, tmp_path, f"m{i}", ["pdf", "docx"])
        assert Path(saida["pdf"]).exists() and Path(saida["docx"]).exists()
    assert plano.tema in _texto_pdf(Path(tmp_path / "m0.pdf"))
    r = materials.exemplo_musical(sessao.adaptador(), tmp_path, "escala", {"tonica": "G", "tipo_escala": "maior", "oitava": 4, "nome_base": "sol"})
    assert Path(r["arquivos"]["midi"]).exists() and Path(r["arquivos"]["musicxml"]).exists()
    r2 = materials.exemplo_musical(sessao.adaptador(), tmp_path, "padrao_ritmico", {"figuras": [1, 0.5, 0.5], "compasso": "2/4", "nome_base": "rit"})
    assert Path(r2["arquivos"]["midi"]).exists()
    assert materials.qr_registro(reg, tmp_path).exists()


def test_relatorio_turma_por_aluno(sessao):
    from percurso.core.models import Agenda, AlunoTurma, Aula, ConteudoTrabalhado, FrequenciaTurma, Registro, Turma
    from percurso.core import state

    reg = sessao.repo.criar_registro(Registro(codigo="", tipo="turma", modalidade="percepcao_musical", identificacao="Percepção 2A", instrumento="percepcao_musical", agenda=Agenda(dia_habitual="quarta", inicio="14:00", termino="14:45", duracao_min=45), turma=Turma(nome="Percepção 2A", quantidade_alunos=2, faixa_etaria="9 a 11 anos", alunos=[AlunoTurma(id="A01", identificacao="Ana"), AlunoTurma(id="A02", identificacao="Bruno")])))
    for d, pres in [("2026-08-05", ["A01", "A02"]), ("2026-08-12", ["A01"]), ("2026-08-19", ["A01", "A02"])]:
        sessao.repo.gravar_aula(reg.codigo, Aula(id="", data=d, duracao_real_min=45, frequencia_turma=FrequenciaTurma(presentes=pres, ausentes=[a for a in ["A01", "A02"] if a not in pres]), classificacao_conteudos=[ConteudoTrabalhado(conteudo="ditado rítmico", situacao="em_desenvolvimento")]))
    state.atualizar_apos_aula(sessao.repo, reg.codigo)
    sessao.carregar(reg.codigo)
    rel = builder.montar(sessao, "frequencia", com_ia=False)
    sec = next(s for s in rel.secoes if s.titulo == "Frequência por aluno")
    linhas = {l[0]: l[3] for l in sec.tabelas[0].linhas}
    assert linhas["Ana"] == "100.0 %" and linhas["Bruno"].startswith("66.7")
